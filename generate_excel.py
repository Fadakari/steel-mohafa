import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Data fetched from Prisma
alloys = [
  {"name": "304", "basePrice": 668000},
  {"name": "316", "basePrice": 1109214},
  {"name": "201", "basePrice": 469565},
  {"name": "321", "basePrice": 298000},
  {"name": "430", "basePrice": 338000},
  {"name": "410", "basePrice": 156000},
  {"name": "310", "basePrice": 1350000},
  {"name": "420", "basePrice": 476190},
  {"name": "309", "basePrice": 365000}
]

# Sort by name just for nicer display
alloys.sort(key=lambda x: x["name"])

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "آپدیت قیمت پایه‌ها"
ws.sheet_view.rightToLeft = True

# Colors
header_fill = PatternFill(start_color="84012B", end_color="84012B", fill_type="solid")
alt_row_fill = PatternFill(start_color="F9F9F9", end_color="F9F9F9", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

header_font = Font(name='Tahoma', size=11, bold=True, color="FFFFFF")
data_font = Font(name='Tahoma', size=11, bold=False, color="000000")
grade_font = Font(name='Tahoma', size=12, bold=True, color="000000")

thin_border = Border(left=Side(style='thin', color='DDDDDD'), 
                     right=Side(style='thin', color='DDDDDD'), 
                     top=Side(style='thin', color='DDDDDD'), 
                     bottom=Side(style='thin', color='DDDDDD'))

center_align = Alignment(horizontal="center", vertical="center")

# Define Headers
headers = ["ردیف", "گرید استیل (آلیاژ)", "قیمت پایه فعلی (تومان)"]
for i in range(1, 11):
    headers.append(f"تاریخ آپدیت {i}: .......")

# Write Headers
for col_num, header_text in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num, value=header_text)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center_align
    cell.border = thin_border

# Write Data
for row_num, alloy in enumerate(alloys, 2):
    is_alt = (row_num % 2 == 0)
    fill = alt_row_fill if is_alt else white_fill
    
    # Radif
    c1 = ws.cell(row=row_num, column=1, value=row_num - 1)
    # Grade
    c2 = ws.cell(row=row_num, column=2, value=f"گرید {alloy['name']}")
    # Current Price
    c3 = ws.cell(row=row_num, column=3, value=alloy["basePrice"])
    
    for c in [c1, c2, c3]:
        c.fill = fill
        c.font = grade_font if c.column == 2 else data_font
        c.alignment = center_align
        c.border = thin_border
        if c.column == 3:
            c.number_format = '#,##0'

    # Format empty columns for future updates
    for col_num in range(4, len(headers) + 1):
        c_empty = ws.cell(row=row_num, column=col_num)
        c_empty.fill = fill
        c_empty.font = data_font
        c_empty.alignment = center_align
        c_empty.border = thin_border
        c_empty.number_format = '#,##0'

# Adjust Column Widths
ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 20
ws.column_dimensions['C'].width = 25
for col_num in range(4, len(headers) + 1):
    col_letter = get_column_letter(col_num)
    ws.column_dimensions[col_letter].width = 22

# Freeze Panes: Freeze the first row and the first two columns
ws.freeze_panes = 'C2'

# Save the file to Downloads
save_path = r'C:\Users\ASUS\Downloads\BasePrices_Steel_Mahfa.xlsx'
wb.save(save_path)
print(f"Saved Excel file to {save_path}")
