import pymysql
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Connect to database
try:
    conn = pymysql.connect(host='localhost', user='root', password='', db='steel_mahfa', charset='utf8mb4')
    cursor = conn.cursor(pymysql.cursors.DictCursor)
    
    # Query joined data
    query = """
    SELECT 
        cbp.id, 
        c.title AS category_name, 
        cbp.base_price 
    FROM category_base_prices cbp
    LEFT JOIN categories c ON cbp.category_id = c.id
    ORDER BY c.title ASC
    """
    cursor.execute(query)
    base_prices = cursor.fetchall()
    conn.close()
except Exception as e:
    print("Database error:", e)
    exit(1)

# Create Excel
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "آپدیت قیمت پایه‌ها"
ws.sheet_view.rightToLeft = True

# Colors & Fonts
header_fill = PatternFill(start_color="84012B", end_color="84012B", fill_type="solid")
alt_row_fill = PatternFill(start_color="F9F9F9", end_color="F9F9F9", fill_type="solid")
white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

header_font = Font(name='Tahoma', size=11, bold=True, color="FFFFFF")
data_font = Font(name='Tahoma', size=11, bold=False, color="000000")
grade_font = Font(name='Tahoma', size=12, bold=True, color="000000")

thin_border = Border(left=Side(style='thin', color='DDDDDD'), 
                     right=Side(style='thin', color='DDDDDD'), 
                     top=Side(style='thin', color='DDDDDD'), 
                     bottom=Side(style='thin', color='DDDDDD'))

center_align = Alignment(horizontal="center", vertical="center")
right_align = Alignment(horizontal="right", vertical="center")

# Headers
headers = ["ردیف", "نام دسته‌بندی (محصولات)", "قیمت پایه فعلی (تومان)"]
for i in range(1, 11):
    headers.append(f"تاریخ آپدیت {i}: .......")

for col_num, header_text in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num, value=header_text)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = center_align
    cell.border = thin_border

# Write Data
for row_num, item in enumerate(base_prices, 2):
    is_alt = (row_num % 2 == 0)
    fill = alt_row_fill if is_alt else white_fill
    
    cat_name = item['category_name'] if item['category_name'] else "دسته بندی نامشخص"
    price = int(item['base_price']) if item['base_price'] else 0
    
    c1 = ws.cell(row=row_num, column=1, value=row_num - 1)
    c2 = ws.cell(row=row_num, column=2, value=cat_name)
    c3 = ws.cell(row=row_num, column=3, value=price)
    
    for c in [c1, c2, c3]:
        c.fill = fill
        c.font = grade_font if c.column == 2 else data_font
        c.alignment = right_align if c.column == 2 else center_align
        c.border = thin_border
        if c.column == 3:
            c.number_format = '#,##0'

    # Future columns
    for col_num in range(4, len(headers) + 1):
        c_empty = ws.cell(row=row_num, column=col_num)
        c_empty.fill = fill
        c_empty.font = data_font
        c_empty.alignment = center_align
        c_empty.border = thin_border
        c_empty.number_format = '#,##0'

# Column Widths
ws.column_dimensions['A'].width = 8
ws.column_dimensions['B'].width = 45  # Made wider for category names
ws.column_dimensions['C'].width = 25
for col_num in range(4, len(headers) + 1):
    col_letter = get_column_letter(col_num)
    ws.column_dimensions[col_letter].width = 22

# Freeze Panes
ws.freeze_panes = 'C2'

# Save
save_path = r'C:\Users\ASUS\Downloads\BasePrices_Steel_Mohafa_v4.xlsx'
wb.save(save_path)
print(f"Saved Excel file to {save_path}")
